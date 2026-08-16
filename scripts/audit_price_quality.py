from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_ROOT = PROJECT_ROOT / "src"
if str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))

from dns_watch_parser.normalizers.price import MAX_PRICE_RUB, normalize_price, repair_concatenated_price


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Audit DNS XLSX price quality before matcher/import")
    parser.add_argument("xlsx", nargs="+", help="DNS_watch_ru_*.xlsx file(s) to validate")
    parser.add_argument("--max-price", type=int, default=MAX_PRICE_RUB, help="Maximum allowed price value")
    parser.add_argument("--examples", type=int, default=10, help="How many suspicious examples to print")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    exit_code = 0
    for raw_path in args.xlsx:
        result = audit_xlsx(Path(raw_path), max_price=args.max_price, example_limit=args.examples)
        print(format_result(result))
        if result["suspicious_price_rows"] or result["unrepairable_price_rows"]:
            exit_code = 1
    return exit_code


def audit_xlsx(path: Path, *, max_price: int = MAX_PRICE_RUB, example_limit: int = 10) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    required = {"price", "old_price", "title", "product_url"}
    missing = sorted(required - set(headers))
    if missing:
        workbook.close()
        raise RuntimeError(f"{path}: missing required columns: {', '.join(missing)}")

    index = {name: headers.index(name) for name in headers}
    rows = 0
    rows_without_price = 0
    suspicious_price_rows = 0
    repairable_price_rows = 0
    unrepairable_price_rows = 0
    max_seen_price = 0
    examples: list[dict] = []

    for values in sheet.iter_rows(min_row=2, values_only=True):
        rows += 1
        price = normalize_price(values[index["price"]])
        old_price = normalize_price(values[index["old_price"]])
        if price is None:
            rows_without_price += 1
            continue

        max_seen_price = max(max_seen_price, price)
        if price <= max_price:
            continue

        suspicious_price_rows += 1
        repaired = repair_concatenated_price(price, old_price, max_price=max_price)
        if repaired is None:
            unrepairable_price_rows += 1
        else:
            repairable_price_rows += 1

        if len(examples) < example_limit:
            examples.append(
                {
                    "title": values[index["title"]] or "",
                    "product_url": values[index["product_url"]] or "",
                    "price": price,
                    "old_price": old_price,
                    "repaired_price": repaired,
                }
            )

    workbook.close()
    return {
        "path": str(path),
        "rows": rows,
        "rows_without_price": rows_without_price,
        "max_seen_price": max_seen_price,
        "suspicious_price_rows": suspicious_price_rows,
        "repairable_price_rows": repairable_price_rows,
        "unrepairable_price_rows": unrepairable_price_rows,
        "examples": examples,
    }


def format_result(result: dict) -> str:
    lines = [
        f"File: {result['path']}",
        f"Rows: {result['rows']}",
        f"Rows without price: {result['rows_without_price']}",
        f"Max price: {result['max_seen_price']}",
        f"Suspicious price rows: {result['suspicious_price_rows']}",
        f"Repairable suspicious rows: {result['repairable_price_rows']}",
        f"Unrepairable suspicious rows: {result['unrepairable_price_rows']}",
    ]
    if result["examples"]:
        lines.append("Examples:")
        for example in result["examples"]:
            lines.append(
                "  - "
                f"{example['title']} :: price={example['price']}, "
                f"old_price={example['old_price']}, repaired={example['repaired_price']}, "
                f"url={example['product_url']}"
            )
    return "\n".join(lines)


if __name__ == "__main__":
    raise SystemExit(main())
