from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_ROOT = PROJECT_ROOT / "src"
if str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))

from dns_watch_parser.normalizers.price import MAX_PRICE_RUB, normalize_price


@dataclass(slots=True)
class ExportSummary:
    path: Path
    rows: int
    rows_without_price: int
    suspicious_price_rows: int
    max_price: int


@dataclass(slots=True)
class LogRun:
    started_at: str
    finished_code: str
    total_urls: str
    exported_rows: str
    low_rows: str


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Diagnose recent DNS parser exports and scheduled logs")
    parser.add_argument("--export-dir", default=str(PROJECT_ROOT / "brand_exports"), help="Directory with DNS XLSX files")
    parser.add_argument("--log", default=str(PROJECT_ROOT / "logs" / "scheduled_run.log"), help="Scheduled run log path")
    parser.add_argument("--latest", type=int, default=10, help="How many recent XLSX files to inspect")
    parser.add_argument("--log-runs", type=int, default=5, help="How many recent scheduled runs to print")
    parser.add_argument("--min-rows", type=int, default=1000, help="Minimum healthy exported row count")
    parser.add_argument("--max-price", type=int, default=MAX_PRICE_RUB, help="Maximum healthy price value")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    export_dir = Path(args.export_dir)
    summaries = [summarize_xlsx(path, max_price=args.max_price) for path in recent_exports(export_dir, args.latest)]
    runs = parse_recent_log_runs(Path(args.log), args.log_runs)

    print_export_summaries(summaries, min_rows=args.min_rows)
    print_log_runs(runs)

    if not summaries:
        print("Health: failed, no DNS_watch_ru_*.xlsx files found")
        return 1

    latest = summaries[0]
    failed = latest.rows < args.min_rows or latest.suspicious_price_rows > 0
    if failed:
        print("Health: failed")
        return 1
    print("Health: ok")
    return 0


def recent_exports(export_dir: Path, limit: int) -> list[Path]:
    if not export_dir.exists():
        return []
    files = sorted(export_dir.glob("DNS_watch_ru_*.xlsx"), key=lambda path: path.stat().st_mtime, reverse=True)
    return files[: max(0, limit)]


def summarize_xlsx(path: Path, *, max_price: int) -> ExportSummary:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    index = {name: headers.index(name) for name in headers}

    rows = 0
    rows_without_price = 0
    suspicious_price_rows = 0
    max_seen_price = 0

    price_index = index.get("price")
    for values in sheet.iter_rows(min_row=2, values_only=True):
        rows += 1
        if price_index is None:
            continue
        price = normalize_price(values[price_index])
        if price is None:
            rows_without_price += 1
            continue
        max_seen_price = max(max_seen_price, price)
        if price > max_price:
            suspicious_price_rows += 1

    workbook.close()
    return ExportSummary(
        path=path,
        rows=rows,
        rows_without_price=rows_without_price,
        suspicious_price_rows=suspicious_price_rows,
        max_price=max_seen_price,
    )


def parse_recent_log_runs(path: Path, limit: int) -> list[LogRun]:
    if not path.exists():
        return []

    runs: list[LogRun] = []
    current: dict[str, str] | None = None
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        if "DNS parser scheduled run started" in line:
            if current is not None:
                runs.append(_log_run_from_dict(current))
            current = {"started_at": _line_timestamp(line)}
            continue
        if current is None:
            continue
        if "Exported rows below minimum" in line:
            current["low_rows"] = line.strip()
        elif line.startswith("Total URLs:"):
            current["total_urls"] = _line_value(line)
        elif line.startswith("Exported rows:"):
            current["exported_rows"] = _line_value(line)
        elif "DNS parser scheduled run finished with code" in line or "DNS parser exited with code" in line:
            current["finished_code"] = _line_value(line)

    if current is not None:
        runs.append(_log_run_from_dict(current))
    return runs[-max(0, limit) :]


def _log_run_from_dict(raw: dict[str, str]) -> LogRun:
    return LogRun(
        started_at=raw.get("started_at", ""),
        finished_code=raw.get("finished_code", ""),
        total_urls=raw.get("total_urls", ""),
        exported_rows=raw.get("exported_rows", ""),
        low_rows=raw.get("low_rows", ""),
    )


def _line_timestamp(line: str) -> str:
    match = re.search(r"\[(.*?)\]", line)
    return match.group(1) if match else ""


def _line_value(line: str) -> str:
    return line.rsplit(" ", 1)[-1].strip()


def print_export_summaries(summaries: list[ExportSummary], *, min_rows: int) -> None:
    print("Recent XLSX exports:")
    if not summaries:
        print("  none")
        return

    for item in summaries:
        status = "ok" if item.rows >= min_rows and item.suspicious_price_rows == 0 else "check"
        modified = datetime.fromtimestamp(item.path.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")
        print(
            f"  {item.path.name}: rows={item.rows}, size={item.path.stat().st_size}, "
            f"without_price={item.rows_without_price}, max_price={item.max_price}, "
            f"suspicious_price={item.suspicious_price_rows}, modified={modified}, status={status}"
        )


def print_log_runs(runs: list[LogRun]) -> None:
    print("Recent scheduled runs:")
    if not runs:
        print("  none")
        return

    for run in runs:
        details = [
            f"started={run.started_at or 'unknown'}",
            f"code={run.finished_code or 'unknown'}",
            f"total={run.total_urls or 'unknown'}",
            f"exported={run.exported_rows or 'unknown'}",
        ]
        if run.low_rows:
            details.append("low_rows=yes")
        print("  " + ", ".join(details))


if __name__ == "__main__":
    raise SystemExit(main())
