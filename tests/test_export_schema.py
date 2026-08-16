from openpyxl import load_workbook

from dns_watch_parser.models import OUTPUT_COLUMNS, ProductRecord
from dns_watch_parser.storage.exports import StreamingXlsxWriter


def test_export_schema_columns(tmp_path):
    path = tmp_path / "out.xlsx"
    writer = StreamingXlsxWriter(path)
    writer.append(ProductRecord(title="Смарт-часы Apple Watch", product_url="https://example.test"))
    assert not path.exists()
    assert writer.temp_path.exists()
    writer.close()

    workbook = load_workbook(path)
    sheet = workbook.active
    columns = [cell.value for cell in sheet[1]]

    assert columns == OUTPUT_COLUMNS
    assert sheet.max_row == 2


def test_export_repairs_concatenated_price_before_save(tmp_path):
    path = tmp_path / "brand_exports" / "out.xlsx"
    writer = StreamingXlsxWriter(path)
    writer.append(
        ProductRecord(
            title="Смарт-часы Xiaomi Watch S4 41 mm",
            product_url="https://example.test/xiaomi",
            price=1599919599,
            old_price=19599,
        )
    )
    writer.close()

    workbook = load_workbook(path)
    sheet = workbook.active
    columns = [cell.value for cell in sheet[1]]
    price_index = columns.index("price") + 1
    old_price_index = columns.index("old_price") + 1

    assert sheet.cell(row=2, column=price_index).value == 15999
    assert sheet.cell(row=2, column=old_price_index).value == 19599
    assert writer.price_stats.suspicious_price_rows == 1
    assert writer.price_stats.repaired_price_rows == 1


def test_export_skips_unrepairable_huge_price(tmp_path):
    path = tmp_path / "brand_exports" / "out.xlsx"
    writer = StreamingXlsxWriter(path)
    appended = writer.append(
        ProductRecord(
            title="Suspicious watch",
            product_url="https://example.test/bad",
            price=1599919599,
            old_price=34499,
        )
    )
    writer.close()

    workbook = load_workbook(path)
    sheet = workbook.active

    assert appended is False
    assert sheet.max_row == 1
    assert writer.price_stats.suspicious_price_rows == 1
    assert writer.price_stats.skipped_price_rows == 1
    assert writer.debug_path.exists()


def test_export_can_discard_incomplete_temp_file(tmp_path):
    path = tmp_path / "out.xlsx"
    path.write_text("previous export", encoding="utf-8")

    writer = StreamingXlsxWriter(path)
    writer.append(ProductRecord(title="DNS Watch", product_url="https://example.test/watch", price=9999))
    temp_path = writer.temp_path
    writer.close(commit=False)

    assert path.read_text(encoding="utf-8") == "previous export"
    assert not temp_path.exists()
